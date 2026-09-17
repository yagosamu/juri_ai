"""Blind human review workbook for a random sample of candidates.

Builds a two-sheet Excel workbook for a lawyer who is blind to which sampled candidates were
accepted into the reviewed set. It reads only the fixed calibration sample, the candidate pool,
the triage file (for competitor articles) and the corpus text. It never reads the file where
verdicts are recorded, the file listing which candidates were accepted, or the published agreement
report, and no cell of the workbook contains a candidate id, a category, a verdict or anything
about inclusion.

Usage: .venv/Scripts/python.exe -m evals.groundtruth.golden.human_review_sheet
"""
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

SHUFFLE_SEED = 20260917
COMPETITORS_PER_ITEM = 5
ARTICLE_MAX_CHARS = 6000
COMPETITOR_MAX_CHARS = 1200

REVIEW_HEADER = ["item", "pergunta", "artigo", "texto do artigo", "artigos parecidos",
                 "responde?", "quais parecidos respondem", "pergunta clara?", "comentário"]

TITLE = "Instruções para a revisão"

INTRO_PARAGRAPHS = [
    "Este arquivo faz parte de uma conferência de qualidade de um conjunto de perguntas sobre o "
    "CPC, o CDC, a CLT e a LGPD, usadas para testar um sistema de busca jurídica. A ideia é "
    "comparar, depois, o quanto a sua leitura como advogado bate com o que foi montado para o teste.",

    "Para cada linha da aba \"Revisão\": leia a pergunta, leia o texto do artigo indicado ao lado "
    "e, se quiser, os artigos parecidos que aparecem na mesma linha. Depois preencha as três "
    "colunas de resposta descritas mais abaixo.",

    "Não existe resposta certa esperada aqui. Se a sua leitura discordar do que a planilha sugere, "
    "essa discordância também é um resultado útil para nós — responda o que você realmente pensa, "
    "sem tentar adivinhar o que \"deveria\" ser.",

    "A ordem das linhas foi embaralhada e não segue nenhuma ordem original. Isso é proposital: "
    "você não recebe nenhuma pista sobre o que o sistema concluiu para cada pergunta, para que a "
    "sua avaliação seja independente.",

    "Suas respostas serão publicadas apenas como números agregados de concordância (por exemplo, "
    "\"em X% das linhas o advogado concordou com Y\"). Você será citado como \"um advogado com "
    "registro ativo na OAB\", a menos que prefira que o seu nome apareça — nesse caso, é só avisar.",

    "Cada linha leva, em média, de 3 a 4 minutos para revisar. São 20 linhas ao todo, o que dá "
    "pouco mais de uma hora de trabalho.",

    "Se alguma linha parecer quebrada, incompleta ou sem sentido, avise a pessoa que te enviou "
    "este arquivo antes de responder.",
]

QUESTIONS_TITLE = "As três perguntas de cada linha"

QUESTIONS = [
    ("1. \"O artigo responde a pergunta por completo?\"",
     "Responda \"sim\", \"em parte\" ou \"não\". Marque \"sim\" somente quando aquele artigo, "
     "sozinho, já responde a pergunta por completo, sem precisar de nenhum outro dispositivo."),

    ("2. \"Algum artigo parecido também responde por completo?\"",
     "Escreva os rótulos dos artigos parecidos que também respondem sozinhos, por exemplo "
     "\"C1, C3\", ou escreva \"nenhum\". Os artigos parecidos estão listados na mesma linha, na "
     "coluna \"artigos parecidos\"."),

    ("3. \"A pergunta está clara?\"",
     "Responda \"sim\" ou \"não\". Marque \"não\" quando a pergunta não puder ser entendida "
     "sozinha — por exemplo, se ela se refere a \"esse tipo de processo\" sem dizer qual processo "
     "é esse."),
]

COMMENT_NOTE = ("Há também uma coluna de comentário, livre e opcional, para qualquer observação "
                "que não caiba nas três respostas acima.")


def _truncate(text: str, max_chars: int) -> str:
    """Cut text at max_chars and mark the cut with ' [...]'; text at or under the cap is untouched."""
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + " [...]"


def _span_text(doc_id: str, start, end, corpus: dict[str, str], context: str) -> str:
    if doc_id not in corpus:
        raise ValueError(f"unknown doc_id {doc_id!r} for {context}")
    doc = corpus[doc_id]
    if not isinstance(start, int) or not isinstance(end, int) or not (0 <= start < end <= len(doc)):
        raise ValueError(f"empty or invalid span [{start!r}:{end!r}) for {context}")
    return doc[start:end]


def build_rows(sample_ids: list[str], candidates: list[dict], triage: dict[str, dict],
              corpus: dict[str, str]) -> list[dict]:
    """One row per sampled candidate, in a deterministic shuffled order, blind to golden-set status.

    Order: sample_ids sorted, then random.Random(SHUFFLE_SEED).shuffle, so the row order is fixed
    and unrelated to any input order (in particular, unrelated to whether a candidate ended up in
    the golden set). Item labels item-01..item-NN are assigned after shuffling. Each row carries the
    question, the article reference and text, and up to COMPETITORS_PER_ITEM competitor articles as
    (label, article_ref, text) tuples, taken in triage order.

    Raises ValueError when a sampled id is missing from candidates, has no triage row, has an empty
    or invalid span, or points at a doc id absent from the corpus; the same check applies to every
    competitor span.
    """
    by_id = {c["candidate_id"]: c for c in candidates}
    order = sorted(sample_ids)
    random.Random(SHUFFLE_SEED).shuffle(order)
    rows = []
    for i, candidate_id in enumerate(order, start=1):
        cand = by_id.get(candidate_id)
        if cand is None:
            raise ValueError(f"candidate {candidate_id!r} is not in candidates.jsonl")
        trow = triage.get(candidate_id)
        if trow is None:
            raise ValueError(f"candidate {candidate_id!r} has no triage row")
        article_text = _truncate(
            _span_text(cand["doc_id"], cand["start"], cand["end"], corpus, f"candidate {candidate_id!r}"),
            ARTICLE_MAX_CHARS)
        competitors = []
        for comp in trow.get("competitors", [])[:COMPETITORS_PER_ITEM]:
            text = _truncate(
                _span_text(comp["doc_id"], comp["start"], comp["end"], corpus,
                          f"competitor {comp.get('label')!r} of {candidate_id!r}"),
                COMPETITOR_MAX_CHARS)
            competitors.append((comp["label"], f"{comp['doc_id'].upper()} {comp['header']}", text))
        rows.append({
            "item": f"item-{i:02d}",
            "candidate_id": candidate_id,
            "question": cand["question"],
            "article_ref": f"{cand['doc_id'].upper()} {cand['header']}",
            "article_text": article_text,
            "competitors": competitors,
        })
    return rows


def sheet_order(rows: list[dict]) -> list[dict]:
    """The item-to-candidate mapping, in item order, and nothing else: safe to commit and to audit
    a run against later, since it says nothing about which items were in the golden set."""
    return [{"item": r["item"], "candidate_id": r["candidate_id"]} for r in rows]


def _write_instructions(ws) -> None:
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 100
    row = 1
    ws.cell(row=row, column=1, value=TITLE).font = Font(bold=True, size=14)
    row += 2
    for paragraph in INTRO_PARAGRAPHS:
        cell = ws.cell(row=row, column=1, value=paragraph)
        cell.alignment = wrap
        row += 2
    ws.cell(row=row, column=1, value=QUESTIONS_TITLE).font = Font(bold=True, size=12)
    row += 2
    for question, explanation in QUESTIONS:
        q_cell = ws.cell(row=row, column=1, value=question)
        q_cell.font = Font(bold=True)
        q_cell.alignment = wrap
        row += 1
        e_cell = ws.cell(row=row, column=1, value=explanation)
        e_cell.alignment = wrap
        row += 2
    ws.cell(row=row, column=1, value=COMMENT_NOTE).alignment = wrap


def _write_review_sheet(ws, rows: list[dict]) -> None:
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.append(REVIEW_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    widths = {"A": 10, "B": 45, "C": 16, "D": 55, "E": 55, "F": 14, "G": 26, "H": 15, "I": 30}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for r in rows:
        competitors_cell = "\n\n".join(
            f"[{label}] {ref}: {text}" for label, ref, text in r["competitors"]) or "nenhum"
        ws.append([r["item"], r["question"], r["article_ref"], r["article_text"], competitors_cell,
                   None, None, None, None])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = wrap
    last_row = ws.max_row
    if last_row >= 2:
        answerable = DataValidation(type="list", formula1='"sim,em parte,não"', allow_blank=True)
        ws.add_data_validation(answerable)
        answerable.add(f"F2:F{last_row}")
        clear = DataValidation(type="list", formula1='"sim,não"', allow_blank=True)
        ws.add_data_validation(clear)
        clear.add(f"H2:H{last_row}")


def write_workbook(rows: list[dict], path: Path) -> None:
    """Write the two-sheet review workbook to path, creating its parent directory.

    Sheet 1, "Instruções", is a plain-language guide in Brazilian Portuguese. Sheet 2, "Revisão",
    has one row per item with a frozen header, wrapped text and dropdown validation on the two
    fixed-choice answer columns. Neither sheet ever writes a candidate id, a category, a judge
    verdict or anything about golden-set inclusion.
    """
    wb = Workbook()
    instructions_ws = wb.active
    instructions_ws.title = "Instruções"
    _write_instructions(instructions_ws)
    review_ws = wb.create_sheet("Revisão")
    _write_review_sheet(review_ws, rows)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _load_candidates(path: Path) -> list[dict]:
    """candidates.jsonl as a list of dicts, keeping only the first line seen for each candidate_id."""
    from evals.groundtruth.golden.jsonl_io import read_jsonl_rows

    seen: dict[str, dict] = {}
    for row in read_jsonl_rows(path):
        seen.setdefault(row["candidate_id"], row)
    return list(seen.values())


def main() -> None:
    from evals.groundtruth.config import CALIBRATION_SAMPLE, CANDIDATES, GROUNDTRUTH_DIR, TRIAGE, load_corpus
    from evals.groundtruth.golden.triage import load_triage

    sample = json.loads(CALIBRATION_SAMPLE.read_text(encoding="utf-8"))
    rows = build_rows(sample["candidate_ids"], _load_candidates(CANDIDATES), load_triage(TRIAGE), load_corpus())

    workbook_path = GROUNDTRUTH_DIR / "runtime" / "human_review" / "golden_review.xlsx"
    write_workbook(rows, workbook_path)

    order_path = GROUNDTRUTH_DIR / "golden" / "human_review_order.json"
    order_path.write_text(
        json.dumps({"seed": SHUFFLE_SEED, "order": sheet_order(rows)}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8")

    print(workbook_path)
    print(order_path)
    print(len(rows))


if __name__ == "__main__":
    main()
