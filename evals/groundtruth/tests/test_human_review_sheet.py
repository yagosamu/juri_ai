import inspect
import random

import openpyxl
import pytest

from evals.groundtruth.golden import human_review_sheet as hrs
from evals.groundtruth.golden.human_review_sheet import (
    COMPETITORS_PER_ITEM, SHUFFLE_SEED, build_rows, sheet_order, write_workbook,
)


def _candidate(cid, doc_id, header, question, start, end):
    return {"candidate_id": cid, "doc_id": doc_id, "start": start, "end": end, "header": header, "question": question}


def _competitor(label, doc_id, header, start, end):
    return {"label": label, "doc_id": doc_id, "header": header, "start": start, "end": end}


def _triage_row(cid, competitors):
    return {"candidate_id": cid, "competitors": competitors}


def _ids(n):
    return [f"r1-cpc-{i:03d}" for i in range(n)]


def _uniform_fixture(n, span=100, n_competitors=1):
    """n candidates all pointing at the same tiny corpus, each with n_competitors competitors."""
    corpus = {"cpc": ("artigo principal de teste " * 20)[:span], "cdc": ("artigo concorrente de teste " * 20)[:span]}
    ids = _ids(n)
    candidates = [_candidate(cid, "cpc", "Art. 1", f"pergunta numero {n}?", 0, span) for n, cid in enumerate(ids)]
    triage = {
        cid: _triage_row(cid, [_competitor(f"C{k}", "cdc", f"Art. {k}", 0, span) for k in range(1, n_competitors + 1)])
        for cid in ids
    }
    return ids, candidates, triage, corpus


def test_build_rows_shuffles_deterministically_for_the_seed():
    ids, candidates, triage, corpus = _uniform_fixture(10)
    rows = build_rows(list(reversed(ids)), candidates, triage, corpus)

    expected = sorted(ids)
    random.Random(SHUFFLE_SEED).shuffle(expected)
    assert [r["candidate_id"] for r in rows] == expected

    other = sorted(ids)
    random.Random(SHUFFLE_SEED + 1).shuffle(other)
    assert other != expected


def test_item_labels_are_zero_padded_in_row_order():
    ids, candidates, triage, corpus = _uniform_fixture(20)
    rows = build_rows(ids, candidates, triage, corpus)
    assert [r["item"] for r in rows] == [f"item-{i:02d}" for i in range(1, 21)]


def test_truncation_marks_long_text_and_leaves_short_text_untouched(monkeypatch):
    monkeypatch.setattr(hrs, "ARTICLE_MAX_CHARS", 10)
    monkeypatch.setattr(hrs, "COMPETITOR_MAX_CHARS", 5)
    corpus = {"cpc": "0123456789ABCDEF", "cdc": "abcdefghij"}

    long_candidates = [_candidate("id-long", "cpc", "Art. 1", "pergunta?", 0, 16)]
    long_triage = {"id-long": _triage_row("id-long", [_competitor("C1", "cdc", "Art. 2", 0, 10)])}
    long_row = build_rows(["id-long"], long_candidates, long_triage, corpus)[0]
    assert long_row["article_text"] == "0123456789 [...]"
    assert long_row["competitors"][0][2] == "abcde [...]"

    short_candidates = [_candidate("id-short", "cpc", "Art. 1", "pergunta?", 0, 5)]
    short_triage = {"id-short": _triage_row("id-short", [_competitor("C1", "cdc", "Art. 2", 0, 3)])}
    short_row = build_rows(["id-short"], short_candidates, short_triage, corpus)[0]
    assert short_row["article_text"] == "01234"
    assert short_row["competitors"][0][2] == "abc"


def test_at_most_five_competitors_in_triage_order_keeping_labels():
    corpus = {"cpc": "artigo principal " * 5, "cdc": "artigo concorrente " * 5}
    candidates = [_candidate("id1", "cpc", "Art. 1", "pergunta?", 0, 20)]
    competitors = [_competitor(f"C{i}", "cdc", f"Art. {i}", 0, 5) for i in range(1, 8)]
    triage = {"id1": _triage_row("id1", competitors)}
    row = build_rows(["id1"], candidates, triage, corpus)[0]
    assert len(row["competitors"]) == COMPETITORS_PER_ITEM == 5
    assert [c[0] for c in row["competitors"]] == ["C1", "C2", "C3", "C4", "C5"]


def test_raises_on_missing_candidate():
    with pytest.raises(ValueError):
        build_rows(["ghost"], [], {}, {"cpc": "texto"})


def test_raises_on_missing_triage_row():
    candidates = [_candidate("id1", "cpc", "Art. 1", "pergunta?", 0, 5)]
    corpus = {"cpc": "0123456789"}
    with pytest.raises(ValueError):
        build_rows(["id1"], candidates, {}, corpus)


def test_raises_on_empty_span():
    candidates = [_candidate("id1", "cpc", "Art. 1", "pergunta?", 5, 5)]
    corpus = {"cpc": "0123456789"}
    triage = {"id1": _triage_row("id1", [])}
    with pytest.raises(ValueError):
        build_rows(["id1"], candidates, triage, corpus)


def test_raises_on_unknown_doc_id():
    candidates = [_candidate("id1", "ghost-doc", "Art. 1", "pergunta?", 0, 5)]
    corpus = {"cpc": "0123456789"}
    triage = {"id1": _triage_row("id1", [])}
    with pytest.raises(ValueError):
        build_rows(["id1"], candidates, triage, corpus)


def test_sheet_order_maps_every_item_to_its_candidate_id_and_nothing_else():
    rows = [
        {"item": "item-01", "candidate_id": "id1", "question": "q1", "extra": "x"},
        {"item": "item-02", "candidate_id": "id2", "question": "q2", "extra": "y"},
    ]
    order = sheet_order(rows)
    assert order == [{"item": "item-01", "candidate_id": "id1"}, {"item": "item-02", "candidate_id": "id2"}]
    assert all(set(entry) == {"item", "candidate_id"} for entry in order)


def test_write_workbook_has_two_sheets_header_and_expected_row_count(tmp_path):
    ids, candidates, triage, corpus = _uniform_fixture(20)
    rows = build_rows(ids, candidates, triage, corpus)
    path = tmp_path / "golden_review.xlsx"
    write_workbook(rows, path)

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Instruções", "Revisão"]
    ws = wb["Revisão"]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == ["item", "pergunta", "artigo", "texto do artigo", "artigos parecidos",
                       "responde?", "quais parecidos respondem", "pergunta clara?", "comentário"]
    assert ws.max_row == 21


def test_write_workbook_never_writes_a_candidate_id_anywhere(tmp_path):
    ids, candidates, triage, corpus = _uniform_fixture(5, n_competitors=3)
    rows = build_rows(ids, candidates, triage, corpus)
    path = tmp_path / "golden_review.xlsx"
    write_workbook(rows, path)

    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = "" if cell.value is None else str(cell.value)
                for cid in ids:
                    assert cid not in value, f"candidate id {cid!r} leaked into {ws.title}!{cell.coordinate}"


def test_module_source_never_mentions_forbidden_golden_files():
    source = inspect.getsource(hrs)
    for forbidden in ("judgments", "golden_set", "golden_consensus"):
        assert forbidden not in source


@pytest.mark.parametrize("text, header, expected", [
    ("Art. 1.016. O agravo de instrumento", "Art. 1", "Art. 1.016"),
    ("Art. 1.052 Enquanto não for editada", "Art. 1", "Art. 1.052"),
    ("Art. 5º Aquele que de qualquer forma", "Art. 5", "Art. 5º"),
    ("Art. 457-A. Texto", "Art. 457", "Art. 457-A"),
    ("texto sem cabeçalho de artigo", "Art. 12", "Art. 12"),
])
def test_article_ref_reads_the_article_number_from_the_span_text(text, header, expected):
    assert hrs.article_ref(header, text) == expected


def test_build_rows_uses_the_full_article_number_for_article_and_competitors():
    corpus = {"cpc": "Art. 1.016. Caberá agravo de instrumento. Art. 1.017. A petição de agravo."}
    candidates = [_candidate("r1-cpc-000", "cpc", "Art. 1", "pergunta?", 0, 41)]
    triage = {"r1-cpc-000": _triage_row("r1-cpc-000", [_competitor("C1", "cpc", "Art. 1", 42, len(corpus["cpc"]))])}
    (row,) = build_rows(["r1-cpc-000"], candidates, triage, corpus)
    assert row["article_ref"] == "CPC Art. 1.016"
    assert row["competitors"][0][1] == "CPC Art. 1.017"


def test_instruction_text_has_no_em_or_en_dash():
    texts = list(hrs.INTRO_PARAGRAPHS) + [t for pair in hrs.QUESTIONS for t in pair] + [hrs.COMMENT_NOTE, hrs.TITLE]
    assert not any("—" in t or "–" in t for t in texts)


def test_write_workbook_validates_answers_freezes_header_and_keeps_ids_out_of_properties(tmp_path):
    ids, candidates, triage, corpus = _uniform_fixture(20)
    path = tmp_path / "review.xlsx"
    write_workbook(build_rows(ids, candidates, triage, corpus), path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Revisão"]
    assert ws.freeze_panes == "A2"
    by_range = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    assert set(by_range) == {"F2:F21", "H2:H21"}
    assert by_range["F2:F21"].formula1 == '"sim,em parte,não"'
    assert by_range["H2:H21"].formula1 == '"sim,não"'
    assert all(dv.showErrorMessage and dv.error for dv in by_range.values())
    props = wb.properties
    for name in ("title", "subject", "creator", "description", "keywords"):
        assert "r1-" not in str(getattr(props, name) or "")
